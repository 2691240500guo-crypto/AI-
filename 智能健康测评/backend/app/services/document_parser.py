"""多模态知识库解析服务：
  1. 文本类（.txt/.md）   —— 直接读取
  2. 文档类（.pdf/.docx） —— python-docx / pypdf 解析
  3. 图片类（.jpg/.png等）—— SiliconFlow DeepSeek-OCR / VL 模型 OCR 识别
统一输出: text（全文）+ chunks（分块列表）
"""
import logging
import os
import re
from typing import List, Tuple

from app.services.llm_client import get_llm

logger = logging.getLogger("parser")

# 文件类型分组
TEXT_EXTS = {".txt", ".md", ".log", ".csv", ".json"}
DOC_EXTS = {".pdf", ".docx"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif"}

# 图片 mime 推断
MIME_MAP = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".bmp": "image/bmp", ".webp": "image/webp", ".gif": "image/gif",
}


def guess_type(ext: str) -> str:
    ext = (ext or "").lower()
    if ext in TEXT_EXTS:
        return "text"
    if ext in DOC_EXTS:
        return "doc"
    if ext in IMAGE_EXTS:
        return "image"
    return "other"


def split_chunks(text: str, chunk_size: int = 500,
                 overlap: int = 80) -> List[str]:
    """按字符长度切分文本块（带重叠），保留语义完整性"""
    text = re.sub(r"\n{3,}", "\n\n", text or "").strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        # 尽量在句子边界处切断
        if end < len(text):
            cut = text.rfind("。", start + chunk_size // 2, end)
            if cut == -1:
                cut = text.rfind("\n", start + chunk_size // 2, end)
            if cut != -1:
                end = cut + 1
        chunks.append(text[start:end].strip())
        # 已切到文本末尾则停止，避免尾部重复小切片
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


def parse_bytes(filename: str, data: bytes) -> Tuple[str, str, List[str]]:
    """解析上传文件字节流。

    Returns:
        (content_text, parse_msg, chunks)
    """
    ext = os.path.splitext(filename)[1].lower()
    ftype = guess_type(ext)

    try:
        if ftype == "text":
            content = _parse_text(data)
        elif ftype == "doc":
            content = _parse_doc(filename, data)
        elif ftype == "image":
            content = _parse_image(data, ext)
        else:
            return "", f"不支持的文件类型: {ext}", []

        content = (content or "").strip()
        if not content:
            return "", "未解析出有效内容", []

        chunks = split_chunks(content)
        msg = (f"解析成功：{ftype} 类型，全文 {len(content)} 字，"
               f"切分 {len(chunks)} 块")
        return content, msg, chunks
    except Exception as e:  # noqa: BLE001
        logger.exception("解析失败 %s", filename)
        return "", f"解析失败: {e}", []


def _parse_text(data: bytes) -> str:
    for enc in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_doc(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".docx":
        import io
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        # 表格内容也抽取
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if ext == ".pdf":
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                parts.append(t)
        return "\n".join(parts)

    return ""


def _parse_image(data: bytes, ext: str) -> str:
    """图片 OCR：优先 DeepSeek-OCR，失败回退 Qwen3-VL"""
    mime = MIME_MAP.get(ext, "image/jpeg")
    llm = get_llm()
    return llm.ocr_image(data, mime)
